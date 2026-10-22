"""Sample Devanagari items and build one annotation workbook per writer slot (W1-W6).

Design: each unit is written by WRITERS_PER_UNIT of the N_WRITERS slots. Units are spread over all
C(6,3) = 20 slot triples, so every writer does half the units and every pair of writers shares a fifth.
Writers see only Devanagari text (never the existing romanization or the gold answer), and each writer
gets their units in a different random order. A hidden manifest maps every row back to its source item,
the existing romanization (R-std), and the answer.

Usage: python scripts/prepare_annotation.py
"""
import json
import random
from itertools import combinations
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / 'data' / 'raw'
OUT = ROOT / 'data' / 'annotation'

SEED = 35
N_WRITERS = 6
WRITERS_PER_UNIT = 3
NECCO_PER_CATEGORY = 75      # 6 categories -> 450 items
BELEBELE_QUESTIONS = 150     # whole passages are sampled until at least this many questions
TRIPLES = list(combinations(range(1, N_WRITERS + 1), WRITERS_PER_UNIT))
N_REPEATS = len(TRIPLES)     # one repeated NeCCo unit per triple -> 10 repeats per writer

NECCO_FILES = [  # (Devanagari file, Romanized file)
    ('Nepali-Cultural-Etiqu.csv', 'Romanized-Cultural-Etiqu.csv'),
    ('Nepali-Tukka.csv', 'Romanized-Tukka.csv'),
    ('Nepali_agro.csv', 'Romanized_agro.csv'),
    ('kinship__Nepali.csv', 'Kinship_romanized.csv'),
]
CATEGORY_SLUG = {
    'Cultural Etiquette': 'etiquette', 'उखान/टुक्का': 'proverbs', 'Geography': 'geography',
    'Festivals': 'festivals', 'Agriculture': 'agriculture', 'Kinship & Social Hierarchy': 'kinship',
}
FIELDS = ['question', 'A', 'B', 'C', 'D']


def load_necco():
    rows = []
    for dev_file, rom_file in NECCO_FILES:
        dev = pd.read_csv(RAW / 'necco' / dev_file)
        rom = pd.read_csv(RAW / 'necco' / rom_file).set_index('ID')
        for _, r in dev.iterrows():
            has_rom = r['ID'] in rom.index
            item = {
                'source': 'necco',
                'category': CATEGORY_SLUG[r['Category'].strip()],
                'item_id': f"necco-{CATEGORY_SLUG[r['Category'].strip()]}-{r['ID']}",
                'answer': str(r['Correct Option']).strip(),
                'dev': {'question': r['Question (Nepali)'], **{k: r[k] for k in 'ABCD'}},
                'rstd': None,
            }
            if has_rom:
                q = rom.loc[r['ID']]
                item['rstd'] = {'question': q['Question (Nepali)'], **{k: q[k] for k in 'ABCD'}}
            rows.append(item)
    return rows


def sample_necco(items, rng):
    by_cat = {}
    for it in items:
        by_cat.setdefault(it['category'], []).append(it)
    picked = []
    for cat in sorted(by_cat):
        pool = by_cat[cat]
        # Prefer items that also have an existing romanization, so D / R-std / R-nat can be compared.
        paired = [it for it in pool if it['rstd'] is not None]
        base = paired if len(paired) >= NECCO_PER_CATEGORY else pool
        picked += rng.sample(base, NECCO_PER_CATEGORY)
    return picked


def load_belebele():
    dev = [json.loads(l) for l in open(RAW / 'belebele' / 'npi_Deva.jsonl', encoding='utf-8')]
    lat = [json.loads(l) for l in open(RAW / 'belebele' / 'npi_Latn.jsonl', encoding='utf-8')]
    # The two files are not line-aligned; (link, question_number) is the unique join key.
    by_key = {(l['link'], l['question_number']): l for l in lat}
    lat = [by_key[(d['link'], d['question_number'])] for d in dev]
    assert all(d['correct_answer_num'] == l['correct_answer_num'] for d, l in zip(dev, lat))
    return dev, lat


def sample_belebele(dev, lat, rng):
    passages = {}
    for i, d in enumerate(dev):
        passages.setdefault(d['link'], []).append(i)
    shuffled = sorted(passages)
    rng.shuffle(shuffled)
    links, n_questions = [], 0
    for link in shuffled:
        if n_questions >= BELEBELE_QUESTIONS:
            break
        links.append(link)
        n_questions += len(passages[link])
    units = []
    for p_no, link in enumerate(links, 1):
        idx = passages[link]
        items = []
        for i in idx:
            d, l = dev[i], lat[i]
            items.append({
                'source': 'belebele', 'category': 'reading',
                'item_id': f"belebele-{i}",
                'answer': 'ABCD'[int(d['correct_answer_num']) - 1],
                'dev': {'question': d['question'], **{k: d[f'mc_answer{n}'] for n, k in enumerate('ABCD', 1)}},
                'rstd': {'question': l['question'], **{k: l[f'mc_answer{n}'] for n, k in enumerate('ABCD', 1)}},
            })
        units.append({
            'unit_id': f'belebele-p{p_no:02d}',
            'passage': {'dev': dev[idx[0]]['flores_passage'], 'rstd': lat[idx[0]]['flores_passage']},
            'items': items,
        })
    return units


def unit_rows(unit):
    """Flatten a unit into annotation rows: (row_id, field_label, devanagari, manifest extras)."""
    rows = []
    if unit.get('passage'):
        rows.append((f"{unit['unit_id']}-passage", 'Passage', unit['passage']['dev'],
                     {'item_id': unit['unit_id'], 'field': 'passage', 'r_std': unit['passage']['rstd'], 'answer': ''}))
    for q_no, it in enumerate(unit['items'], 1):
        prefix = f"Q{q_no} " if len(unit['items']) > 1 or unit.get('passage') else ''
        for f in FIELDS:
            label = prefix + ('Question' if f == 'question' else f'Option {f}')
            rid = f"{unit['unit_id']}-{it['item_id']}-{f}" if unit.get('passage') else f"{unit['unit_id']}-{f}"
            rows.append((rid, label, str(it['dev'][f]).strip(),
                         {'item_id': it['item_id'], 'field': f,
                          'r_std': '' if it['rstd'] is None else str(it['rstd'][f]).strip(),
                          'answer': it['answer']}))
    return rows


# ---------------------------------------------------------------- workbook
FONT = 'Arial'
DEV_FONT = 'Nirmala UI'   # Arial has no Devanagari glyphs
HEAD_FILL = PatternFill('solid', start_color='1F3864')
INPUT_FILL = PatternFill('solid', start_color='FFF2CC')
READ_FILL = PatternFill('solid', start_color='F2F2F2')
THIN = Border(*(Side(style='thin', color='BFBFBF'),) * 4)
WRAP = Alignment(wrap_text=True, vertical='top')

INSTRUCTIONS = [
    ('Romanized Nepali writing task', 'title'),
    ('Thank you for helping. Your writing will be used, without your name, to study how AI systems understand '
     'Nepali written in English letters.', None),
    ('What to do', 'head'),
    ('1. Open the "Items" sheet. Each row shows Nepali text in Devanagari.', None),
    ('2. In the yellow column, type the same text in English letters (Roman script), exactly the way YOU would '
     'type it in Messenger, Viber or WhatsApp.', None),
    ('3. Fill in the "About You" sheet once.', None),
    ('Rules', 'head'),
    ('• There is no correct spelling. Write it your own way. Do not try to be consistent with anyone else.', None),
    ('• Keep the same words in the same order. Do not translate to English and do not rephrase.', None),
    ('• Short forms you normally use are fine (for example writing छ as "xa" or "cha").', None),
    ('• Do NOT use Google Translate, a transliteration tool or keyboard, or ask anyone else.', None),
    ('• Some rows may look familiar or repeat. Just write them again normally; do not look back.', None),
    ('• Work in several short sessions if you like; save the file often.', None),
    ('• If a word is unclear, write it as you would guess it and add a note in the Notes column.', None),
    ('Example (two different people, both correct)', 'head'),
    ('Devanagari:  नेपालको राजधानी कुन हो?', 'dev'),
    ('Person 1:  nepal ko rajdhani kun ho?', None),
    ('Person 2:  Nepalko rajdhaani kun ho', None),
    ('Progress', 'head'),
]


def build_workbook(writer, units, path):
    wb = Workbook()
    ins = wb.active
    ins.title = 'Instructions'
    ins.column_dimensions['A'].width = 110
    r = 1
    for text, kind in INSTRUCTIONS:
        if kind == 'head':
            r += 1  # blank line before each heading
        c = ins.cell(row=r, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.font = Font(name=DEV_FONT if kind == 'dev' else FONT, size=16 if kind == 'title' else 11,
                      bold=kind in ('title', 'head'))
        r += 1
    ins.cell(row=r, column=1, value='Rows completed:').font = Font(name=FONT, bold=True)
    ins.cell(row=r + 1, column=1, value="=COUNTA(Items!E:E)-1&\" of \"&(COUNTA(Items!D:D)-1)").font = Font(name=FONT)
    ins.cell(row=r + 3, column=1, value='Legend: yellow cells are the only cells you need to fill in.').font = Font(name=FONT, italic=True)
    ins.cell(row=r + 3, column=1).fill = INPUT_FILL

    about = wb.create_sheet('About You')
    about.column_dimensions['A'].width = 52
    about.column_dimensions['B'].width = 40
    questions = [
        ('Writer code (already filled)', writer),
        ('Home district / region where you grew up', ''),
        ('First language(s)', ''),
        ('Age group (under 20 / 20-29 / 30-39 / 40+)', ''),
        ('School medium (Nepali / English / both)', ''),
        ('How often do you type Nepali in English letters? (daily / weekly / rarely)', ''),
        ('Device you usually type on (phone / computer)', ''),
        ('Do you usually mix English words when you chat? (often / sometimes / rarely)', ''),
        ('I agree that my anonymous writing may be published in a research dataset (yes / no)', ''),
    ]
    for i, (q, v) in enumerate(questions, 1):
        about.cell(row=i, column=1, value=q).font = Font(name=FONT)
        c = about.cell(row=i, column=2, value=v)
        c.font = Font(name=FONT, bold=bool(v))
        if not v:
            c.fill = INPUT_FILL
        c.border = THIN
    dv = DataValidation(type='list', formula1='"yes,no"', allow_blank=True)
    about.add_data_validation(dv)
    dv.add(f'B{len(questions)}')

    items = wb.create_sheet('Items')
    headers = ['Row ID', 'Unit', 'Part', 'Devanagari (read only)', 'Your Romanized writing', 'Notes (optional)']
    widths = [30, 7, 14, 60, 60, 25]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        c = items.cell(row=1, column=col, value=h)
        c.font = Font(name=FONT, bold=True, color='FFFFFF')
        c.fill = HEAD_FILL
        c.alignment = WRAP
        items.column_dimensions[c.column_letter].width = w
    items.freeze_panes = 'D2'
    r = 2
    manifest = []
    for u_no, unit in enumerate(units, 1):
        for rid, label, dev, extra in unit_rows(unit):
            # Writers see an opaque code; the real key (which reveals repeats) stays in the manifest.
            code = f'{writer}-{r - 1:04d}'
            vals = [code, u_no, label, dev, None, None]
            for col, v in enumerate(vals, 1):
                c = items.cell(row=r, column=col, value=v)
                c.alignment = WRAP
                c.border = THIN
                c.font = Font(name=DEV_FONT if col == 4 else FONT, size=11 if col in (4, 5) else 9,
                              color='808080' if col == 1 else None)
                if col == 4:
                    c.fill = READ_FILL
                if col in (5, 6):
                    c.fill = INPUT_FILL
            manifest.append({'writer': writer, 'code': code, 'row_key': rid, 'order': r - 1, 'unit_no': u_no,
                             'unit_id': unit['unit_id'], 'repeat_of': unit.get('repeat_of', ''),
                             'source': unit['source'], 'category': unit['category'],
                             'devanagari': dev, **extra})
            r += 1
    wb.calculation.fullCalcOnLoad = True
    wb.active = wb.sheetnames.index('Instructions')
    wb.save(path)
    return manifest


def build_sample():
    """Return (units, repeats). Shared by the spreadsheet build and the web-app seed export."""
    rng = random.Random(SEED)
    necco = sample_necco(load_necco(), rng)
    dev, lat = load_belebele()
    bele_units = sample_belebele(dev, lat, rng)

    units = [{'unit_id': it['item_id'], 'source': 'necco', 'category': it['category'], 'items': [it]}
             for it in necco]
    for u in bele_units:
        u.update(source='belebele', category='reading')
    units += bele_units
    assign_slots(units, rng)
    # One repeated NeCCo unit per writer triple, so every writer gets the same number of repeats.
    # A repeat goes to the same writers as its original.
    necco_units = [u for u in units if u['source'] == 'necco']
    originals = [rng.choice([u for u in necco_units if tuple(u['writers']) == t]) for t in TRIPLES]
    repeats = [{**u, 'unit_id': f"repeat-{i:02d}", 'repeat_of': u['unit_id']}
               for i, u in enumerate(originals, 1)]
    return units, repeats


def assign_slots(units, rng):
    """Give every unit a writer triple, cycling through all 20 triples within each category in turn,
    so each writer gets a balanced half of every category and of the long Belebele passages."""
    by_cat = {}
    for u in units:
        by_cat.setdefault(u['category'], []).append(u)
    i = 0
    for cat in sorted(by_cat):
        group = by_cat[cat]
        rng.shuffle(group)
        for u in group:
            u['writers'] = list(TRIPLES[i % len(TRIPLES)])
            i += 1


def main():
    units, repeats = build_sample()

    (OUT / 'sheets').mkdir(parents=True, exist_ok=True)
    (OUT / 'key').mkdir(parents=True, exist_ok=True)
    for old in (OUT / 'sheets').glob('romanization_W*.xlsx'):
        old.unlink()
    manifest = []
    for slot in range(1, N_WRITERS + 1):
        writer = f'W{slot}'
        order = [u for u in units if slot in u['writers']]
        random.Random(SEED + slot).shuffle(order)
        # Repeats go in the second half, so they are far from most first appearances.
        half = len(order) // 2
        for rep in (r for r in repeats if slot in r['writers']):
            order.insert(random.Random(f'{SEED}-{writer}-{rep["unit_id"]}').randint(half, len(order)), rep)
        manifest += build_workbook(writer, order, OUT / 'sheets' / f'romanization_{writer}.xlsx')

    m = pd.DataFrame(manifest)
    m.to_csv(OUT / 'key' / 'manifest.csv', index=False, encoding='utf-8')
    n_q = sum(len(u['items']) for u in units)
    print(f"units: {len(units)} (+{N_REPEATS} repeats) | questions: {n_q}")
    m['words'] = m.devanagari.str.split().str.len()
    per = m.groupby('writer').agg(rows=('code', 'size'), words=('words', 'sum'),
                                  repeats=('repeat_of', lambda s: s[s.fillna('') != ''].nunique()))
    print(per.to_string())
    print(m[m.writer == 'W1'].groupby(['source', 'category']).item_id.nunique().to_string())


if __name__ == '__main__':
    main()
